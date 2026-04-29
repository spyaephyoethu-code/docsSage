import io
import logging

import docx
import httpx

# runtime imports
from app.infrastructure.ingestion.base import Document
from app.infrastructure.ingestion.loaders.base import BaseLoader

logger = logging.getLogger(__name__)


class WordLoader(BaseLoader):
    async def load(self, url_or_path: str) -> list[Document]:
        if url_or_path.startswith("http"):
            async with httpx.AsyncClient(timeout=60, follow_redirects=True) as client:
                resp = await client.get(url_or_path)
                resp.raise_for_status()
                doc = docx.Document(io.BytesIO(resp.content))
        else:
            doc = docx.Document(url_or_path)

        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        text = "\n\n".join(paragraphs)

        if not text:
            return []

        logger.info("WordLoader: loaded 1 document")
        return [Document(text=text, metadata={"source_type": "word", "url": url_or_path})]
